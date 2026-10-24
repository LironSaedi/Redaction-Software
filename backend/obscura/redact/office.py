"""Word / PowerPoint / Excel redaction.

DOCX and PPTX are processed at the XML level so *every* text carrier is covered: body, tables, text
boxes, headers, footers, footnotes, endnotes, SmartArt, chart caches and speaker notes. Before analysis
the package is cleaned of the places PII hides out of sight: tracked-change deletions, comments, hidden
(vanish) runs, field codes, hyperlink targets, document properties, custom XML, embedded OLE objects
and thumbnails. XLSX is processed cell by cell with each value read in the context of its column
header ("SSN: 123-45-6789"), which makes label-driven detection work on spreadsheets.
"""
from __future__ import annotations

import io
import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path

from lxml import etree

from ..detection.text.engine import AnalyzeOptions, SearchTerm, TextAnalyzer, get_analyzer, partial_ranges
from ..detection.text.entities import ENTITIES
from .model import Detection

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
NS = {"w": W, "a": A}

_DOCX_TEXT_PARTS = re.compile(r"^word/(document|header\d*|footer\d*|footnotes|endnotes|glossary/document)\.xml$")
_MEDIA = re.compile(r"^(?:word|ppt)/media/.+\.(?:jpe?g|png|tiff?|bmp|gif|webp)$", re.I)
_DOCX_DRAWING_PARTS = re.compile(r"^word/(diagrams|charts)/.*\.xml$")
_PPTX_TEXT_PARTS = re.compile(r"^ppt/(slides/slide\d+|notesSlides/notesSlide\d+|slideLayouts/slideLayout\d+|"
                              r"slideMasters/slideMaster\d+|diagrams/.+|charts/.+)\.xml$")
_DROP_PARTS = re.compile(r"^(?:customXml/.*|docProps/custom\.xml|docProps/thumbnail\.\w+|word/comments.*\.xml|"
                         r"word/people\.xml|word/embeddings/.*|ppt/embeddings/.*|ppt/comments/.*|ppt/commentAuthors\.xml|"
                         r"word/vbaProject\.bin|ppt/vbaProject\.bin)$")


def redaction_text(d: Detection, style: str, original: str) -> str:
    """Replacement for redacted characters in text formats."""
    if style == "block":
        return "".join(" " if c == " " else "█" for c in original)
    if style == "entity":
        e = ENTITIES.get(d.entity)
        return f"[{(e.label if e else d.entity).upper()}]"
    if style == "exemption":
        return f"[{d.exemption or 'REDACTED'}]"
    return "[REDACTED]"


# ---------------------------------------------------------------------------------------------
# OOXML (docx / pptx)
# ---------------------------------------------------------------------------------------------

@dataclass
class _Seg:
    part: str
    para_idx: int
    node_idx: int  # index into that paragraph's text nodes
    start: int     # offset in full text
    end: int


class OoxmlDocument:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.kind = self.path.suffix.lower().lstrip(".")
        if self.kind not in ("docx", "pptx", "docm", "pptm"):
            raise ValueError("not a docx/pptx file")
        self.zin = zipfile.ZipFile(self.path)
        self.trees: dict[str, etree._ElementTree] = {}
        self.paras: dict[str, list[list[etree._Element]]] = {}
        self.text = ""
        self.segs: list[_Seg] = []
        self.cleaned: dict[str, int] = {}
        self._load()

    def _is_text_part(self, name: str) -> bool:
        if self.kind.startswith("docx") or self.kind == "docm":
            return bool(_DOCX_TEXT_PARTS.match(name) or _DOCX_DRAWING_PARTS.match(name))
        return bool(_PPTX_TEXT_PARTS.match(name))

    def _load(self) -> None:
        parts = sorted(n for n in self.zin.namelist() if self._is_text_part(n))
        # document body first, then headers/footers/etc.
        parts.sort(key=lambda n: (0 if n.endswith(("document.xml",)) or "/slides/" in n else 1, _natural(n)))
        buf: list[str] = []
        pos = 0
        for name in parts:
            tree = etree.parse(io.BytesIO(self.zin.read(name)))
            self._clean_tree(tree)
            self.trees[name] = tree
            plist: list[list[etree._Element]] = []
            prefixes = _table_prefixes(tree)
            for p in tree.iter(f"{{{W}}}p", f"{{{A}}}p"):
                nodes = [t for t in p.iter(f"{{{W}}}t", f"{{{A}}}t", f"{{{W}}}tab", f"{{{W}}}br", f"{{{A}}}br")
                         if _nearest_para(t) is p]
                if not nodes:
                    continue
                if p in prefixes:
                    # a table cell is read as "Column header: value" so header labels give context
                    # (the header text is not part of any node, so it is never edited)
                    buf.append(prefixes[p])
                    pos += len(prefixes[p])
                pi = len(plist)
                plist.append(nodes)
                for ni, node in enumerate(nodes):
                    tag = etree.QName(node).localname
                    s = "\t" if tag == "tab" else ("\n" if tag == "br" else (node.text or ""))
                    if tag in ("t",):
                        self.segs.append(_Seg(name, pi, ni, pos, pos + len(s)))
                    buf.append(s)
                    pos += len(s)
                buf.append("\n")
                pos += 1
            if "/charts/" in name:
                # chart caches: series names and category labels are plain <c:v> strings, not runs
                for v in tree.iter(f"{{{C}}}v"):
                    if v.text and any(etree.QName(a).localname in ("strCache", "strRef", "multiLvlStrCache", "tx")
                                      for a in v.iterancestors()):
                        pi = len(plist)
                        plist.append([v])
                        self.segs.append(_Seg(name, pi, 0, pos, pos + len(v.text)))
                        buf.append(v.text + "\n")
                        pos += len(v.text) + 1
            self.paras[name] = plist
            buf.append("\n")
            pos += 1
        self.segs.sort(key=lambda g: g.start)
        self.text = "".join(buf)

    def _clean_tree(self, tree: etree._ElementTree) -> None:
        root = tree.getroot()
        c = self.cleaned
        # tracked changes: deleted text is invisible but still in the file
        for tag in ("del", "moveFrom"):
            for el in list(root.iter(f"{{{W}}}{tag}")):
                el.getparent().remove(el)
                c["tracked_deletions"] = c.get("tracked_deletions", 0) + 1
        for tag in ("ins", "moveTo"):
            for el in list(root.iter(f"{{{W}}}{tag}")):
                _unwrap(el)
        # comments anchors
        for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
            for el in list(root.iter(f"{{{W}}}{tag}")):
                parent = el.getparent()
                parent.remove(el)
                if etree.QName(parent).localname == "r" and not len(parent):
                    parent.getparent().remove(parent)
                c["comments"] = c.get("comments", 0) + 1
        # hidden (vanish) runs
        for run in list(root.iter(f"{{{W}}}r")):
            rpr = run.find(f"{{{W}}}rPr")
            if rpr is not None and rpr.find(f"{{{W}}}vanish") is not None:
                run.getparent().remove(run)
                c["hidden_runs"] = c.get("hidden_runs", 0) + 1
        # field codes (HYPERLINK "mailto:…", MERGEFIELD …) and hyperlink wrappers
        for el in list(root.iter(f"{{{W}}}instrText")):
            el.text = ""
            c["field_codes"] = c.get("field_codes", 0) + 1
        for el in list(root.iter(f"{{{W}}}hyperlink")):
            _unwrap(el)
            c["hyperlinks"] = c.get("hyperlinks", 0) + 1
        for el in list(root.iter(f"{{{A}}}hlinkClick")):
            el.getparent().remove(el)
            c["hyperlinks"] = c.get("hyperlinks", 0) + 1
        # alt text / descriptions on images can name people
        for el in root.iter():
            for attr in ("descr", "title"):
                if el.get(attr):
                    el.set(attr, "")
                    c["alt_text"] = c.get("alt_text", 0) + 1

    # -----------------------------------------------------------------------------------------

    def analyze(self, options: AnalyzeOptions | None = None, analyzer: TextAnalyzer | None = None,
                images: bool = True) -> list[Detection]:
        an = analyzer or get_analyzer()
        opts = options or AnalyzeOptions()
        dets = [_span_to_det(sp) for sp in an.analyze(self.text, opts)]
        if images:
            dets += self.image_detections(opts, an, known=dets)
        return dets

    def image_detections(self, options: AnalyzeOptions, analyzer: TextAnalyzer,
                         known: list[Detection] = ()) -> list[Detection]:
        """Pictures inside the document — screenshots of texts, scanned inserts, photos of cards, chart images —
        are OCR'd and checked for faces / plates. Values already found in the document text (names, phones…)
        are searched for in the pictures as well. Boxes are in image pixels and live in ``meta``; the part is
        redacted at the pixel level on export."""
        from dataclasses import replace

        from ..storage import get_storage
        from .image import ImageDocument
        from ..detection.text.names import NON_NAME_WORDS
        terms = [SearchTerm(d.text, entity=d.entity) for d in known
                 if d.selected and d.entity in ("PHONE", "EMAIL", "SSN", "ADDRESS", "MRN", "DATE_OF_BIRTH")
                 and len(d.text.strip()) >= 4]
        # names appear in other forms in screenshots ("DELGADO-RUIZ, MARIA" in the text, "Maria" in a chat)
        name_tokens = {w.strip(".,") for d in known if d.selected and d.entity in ("PERSON", "MINOR")
                       for w in re.split(r"[\s,]+", d.text) if len(w.strip(".,")) >= 3}
        terms += [SearchTerm(w, entity="PERSON") for w in sorted(name_tokens) if w.lower() not in NON_NAME_WORDS]
        out: list[Detection] = []
        for name in sorted(n for n in self.zin.namelist() if _MEDIA.match(n)):
            data = self.zin.read(name)
            with get_storage().scratch(Path(name).suffix.lower()) as tmp:
                tmp.write_bytes(data)
                try:
                    idoc = ImageDocument(tmp)
                except Exception:
                    continue  # not a raster image we can read (EMF/WMF, corrupt)
                if idoc.w < 60 or idoc.h < 20:
                    continue
                opts = replace(options, terms=list(options.terms) + terms)
                found = idoc.analyze(classes={"face", "license_plate"}, options=opts, analyzer=analyzer)
                found.append(idoc.classification_detection(faces=sum(1 for d in found if d.entity == "FACE")))
            for d in found:
                d.meta = dict(d.meta, media=name, image_boxes=[b.to_dict() for b in d.boxes],
                              image_size=[idoc.w, idoc.h])
                d.boxes = []
                d.start = d.end = -1
                d.explanation = f"in picture {Path(name).name}: " + d.explanation
                out.append(d)
        return out

    def search(self, terms: list[SearchTerm]) -> list[Detection]:
        spans = TextAnalyzer._user_rules(self.text, terms)
        return [_span_to_det(sp) for sp in spans]

    def apply(self, detections: list[Detection], out_path: str | Path, style: str = "block") -> dict:
        ranges: list[tuple[int, int, Detection]] = []
        for d in detections:
            if not d.selected or d.start < 0:
                continue
            if d.partial:
                for a, b in partial_ranges(self.text[d.start:d.end], d.partial):
                    ranges.append((d.start + a, d.start + b, d))
            else:
                ranges.append((d.start, d.end, d))
        ranges.sort()
        # per text node: list of (local_start, local_end, replacement) – replacement only on first node
        edits: dict[tuple[str, int, int], list[tuple[int, int, str]]] = {}
        for a, b, d in ranges:
            first = True
            rep = redaction_text(d, style, self.text[a:b])
            for seg in self._segs_overlapping(a, b):
                ls, le = max(a, seg.start) - seg.start, min(b, seg.end) - seg.start
                key = (seg.part, seg.para_idx, seg.node_idx)
                if style == "block":
                    piece = rep[max(a, seg.start) - a:min(b, seg.end) - a]
                else:
                    piece = rep if first else ""
                edits.setdefault(key, []).append((ls, le, piece))
                first = False
        for (part, pi, ni), lst in edits.items():
            node = self.paras[part][pi][ni]
            s = node.text or ""
            for ls, le, piece in sorted(lst, reverse=True):
                s = s[:ls] + piece + s[le:]
            node.text = s
            if etree.QName(node).localname == "t":
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        out_path = Path(out_path)
        img_report = self._redact_media([d for d in detections if d.selected and d.meta.get("media")])
        self._write(out_path)
        report = {"redactions": len(ranges), "cleaned": self.cleaned, "images": img_report}
        report["verification"] = verify_zip(out_path, [d for d in detections if d.selected and not d.partial
                                                       and not d.meta.get("media")])
        bad = [r for r in img_report.values() if not r.get("verification", {}).get("passed", True)]
        if bad:
            report["verification"]["passed"] = False
            report["verification"]["issues"].append({"type": "image_redaction_unverified", "count": len(bad)})
        return report

    def _redact_media(self, dets: list[Detection]) -> dict:
        """Burn the selected boxes into each affected picture (a new image object: no metadata survives)."""
        from ..storage import get_storage
        from .image import ImageDocument
        from .model import Box
        self._media_override: dict[str, bytes] = getattr(self, "_media_override", {})
        by_part: dict[str, list[Detection]] = {}
        for d in dets:
            by_part.setdefault(d.meta["media"], []).append(d)
        report = {}
        for name, lst in by_part.items():
            suffix = Path(name).suffix.lower()
            st = get_storage()
            with st.scratch(suffix) as src, st.scratch(suffix) as dst:
                src.write_bytes(self.zin.read(name))
                pix = [Detection(entity=d.entity, text=d.text, score=d.score, source=d.source, kind=d.kind,
                                 boxes=[Box.from_dict(b) for b in d.meta.get("image_boxes", [])], selected=True)
                       for d in lst]
                rep = ImageDocument(src).apply(pix, dst, effect="blackout")
                self._media_override[name] = dst.read_bytes()
            report[name] = {"redactions": rep["redactions"], "verification": rep.get("verification", {})}
        return report


    def _segs_overlapping(self, a: int, b: int):
        # segs are sorted by start
        lo, hi = 0, len(self.segs)
        while lo < hi:
            mid = (lo + hi) // 2
            if self.segs[mid].end <= a:
                lo = mid + 1
            else:
                hi = mid
        i = lo
        while i < len(self.segs) and self.segs[i].start < b:
            if self.segs[i].end > a:
                yield self.segs[i]
            i += 1

    def _write(self, out_path: Path) -> None:
        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in self.zin.infolist():
                name = item.filename
                if _DROP_PARTS.match(name):
                    self.cleaned["parts_removed"] = self.cleaned.get("parts_removed", 0) + 1
                    continue
                if name in self.trees:
                    data = etree.tostring(self.trees[name], xml_declaration=True, encoding="UTF-8", standalone=True)
                elif name == "docProps/core.xml":
                    data = _blank_core_props(self.zin.read(name))
                elif name == "docProps/app.xml":
                    data = _blank_app_props(self.zin.read(name))
                elif name.endswith(".rels"):
                    data = _clean_rels(self.zin.read(name), name)
                elif name == "[Content_Types].xml":
                    data = _clean_content_types(self.zin.read(name))
                elif name == "word/settings.xml":
                    data = _clean_settings(self.zin.read(name))
                elif name in getattr(self, "_media_override", {}):
                    data = self._media_override[name]
                elif re.match(r"^(?:word|ppt)/media/.+\.(?:jpe?g|png|tiff?)$", name, re.I):
                    data = _strip_image_metadata(self.zin.read(name), name)
                else:
                    data = self.zin.read(name)
                zout.writestr(item, data)


def _table_prefixes(tree) -> dict:
    """Map the first paragraph of every body cell to "Header: " from the table's first row, when that row
    looks like a header (short cells, mostly words)."""
    out = {}
    for tbl in tree.iter(f"{{{W}}}tbl", f"{{{A}}}tbl"):
        rows = [r for r in tbl if etree.QName(r).localname == "tr"]
        if len(rows) < 2:
            continue

        def cells(row):
            return [c for c in row if etree.QName(c).localname == "tc"]

        def text(cell):
            return " ".join((t.text or "") for t in cell.iter(f"{{{W}}}t", f"{{{A}}}t")).strip()

        header = [text(c) for c in cells(rows[0])]
        if not header or any(len(h) > 40 for h in header) or sum(bool(re.search(r"\d", h)) for h in header) > len(header) / 2:
            continue
        for row in rows[1:]:
            for ci, cell in enumerate(cells(row)):
                if ci < len(header) and header[ci]:
                    first = next(cell.iter(f"{{{W}}}p", f"{{{A}}}p"), None)
                    if first is not None:
                        out[first] = header[ci] + ": "
    return out


def _natural(s: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]


def _nearest_para(el):
    p = el.getparent()
    while p is not None and etree.QName(p).localname != "p":
        p = p.getparent()
    return p


def _unwrap(el) -> None:
    parent = el.getparent()
    idx = parent.index(el)
    for child in list(el):
        parent.insert(idx, child)
        idx += 1
    parent.remove(el)


def _blank_core_props(data: bytes) -> bytes:
    tree = etree.fromstring(data)
    for el in tree:
        if etree.QName(el).localname in ("creator", "lastModifiedBy", "title", "subject", "description", "keywords",
                                         "category", "contentStatus", "identifier", "language", "version"):
            el.text = ""
    return etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)


def _blank_app_props(data: bytes) -> bytes:
    """app.xml keeps "TitlesOfParts" — a copy of every slide / sheet / document title — plus company data."""
    tree = etree.fromstring(data)
    for el in tree.iter():
        if etree.QName(el).localname in ("Company", "Manager", "HyperlinkBase", "Template"):
            el.text = ""
    for top in tree.iter():
        if etree.QName(top).localname == "TitlesOfParts":
            for el in top.iter():
                if etree.QName(el).localname == "lpstr":
                    el.text = ""
    return etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)


def _rel_target(rels_name: str, target: str) -> str:
    """Resolve a relationship target to a package part name."""
    import posixpath
    if target.startswith("/"):
        return target.lstrip("/")
    src_dir = posixpath.dirname(posixpath.dirname(rels_name))  # word/_rels/document.xml.rels -> word
    return posixpath.normpath(posixpath.join(src_dir, target)).lstrip("./") if src_dir else posixpath.normpath(target)


def _clean_rels(data: bytes, rels_name: str = "") -> bytes:
    tree = etree.fromstring(data)
    for rel in list(tree):
        t = rel.get("Type", "")
        external = rel.get("TargetMode") == "External"
        if t == R_HYPERLINK or external:
            rel.set("Target", "#")
            continue
        if _DROP_PARTS.match(_rel_target(rels_name, rel.get("Target", ""))):
            tree.remove(rel)
    return etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_content_types(data: bytes) -> bytes:
    tree = etree.fromstring(data)
    for el in list(tree):
        pn = el.get("PartName", "")
        if pn and _DROP_PARTS.match(pn.lstrip("/")):
            tree.remove(el)
    return etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)


def _clean_settings(data: bytes) -> bytes:
    tree = etree.fromstring(data)
    for el in list(tree.iter(f"{{{W}}}rsids", f"{{{W}}}attachedTemplate", f"{{{W}}}mailMerge",
                             f"{{{W}}}documentProtection")):
        el.getparent().remove(el)
    return etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)


def _strip_image_metadata(data: bytes, name: str) -> bytes:
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(data))
        out = io.BytesIO()
        fmt = im.format or ("PNG" if name.lower().endswith("png") else "JPEG")
        clean = Image.new(im.mode, im.size)
        clean.putdata(list(im.getdata()))
        clean.save(out, fmt, quality=95) if fmt == "JPEG" else clean.save(out, fmt)
        return out.getvalue()
    except Exception:
        return data


def _span_to_det(sp) -> Detection:
    ent = ENTITIES.get(sp.entity)
    return Detection(entity=sp.entity, text=sp.text, score=sp.score, source=sp.source, boxes=[], page=0,
                     start=sp.start, end=sp.end, explanation=sp.explanation,
                     selected=bool(sp.meta.get("selected", True)),
                     exemption=sp.meta.get("exemption") or (ent.exemption if ent else ""),
                     partial=sp.meta.get("partial"), kind="text",
                     meta={k: v for k, v in sp.meta.items() if k in ("sources", "alt_entities")})


# ---------------------------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------------------------

class XlsxDocument:
    """Spreadsheet redaction. Text is built as "Header: value" lines so column labels act as context."""

    def __init__(self, path: str | Path):
        import openpyxl
        self.path = Path(path)
        self.wb = openpyxl.load_workbook(str(path), data_only=False, keep_vba=False)
        self.text = ""
        self.cells: list[tuple[int, int, object]] = []  # (start, end, cell)
        self.cleaned: dict[str, int] = {}
        self._build()

    def _build(self) -> None:
        buf, pos = [], 0
        self.titles: list[tuple[int, int, object]] = []
        for ws in self.wb.worksheets:
            line = f"Sheet: {ws.title}\n"
            self.titles.append((pos + 7, pos + 7 + len(ws.title), ws))
            buf.append(line)
            pos += len(line)
            rows = list(ws.iter_rows())
            header: dict[int, str] = {}
            header_row = None
            for row in rows:
                vals = [c for c in row if c.value is not None and str(c.value).strip()]
                if vals and sum(isinstance(c.value, str) for c in vals) >= max(1, int(0.7 * len(vals))):
                    header = {c.column: str(c.value).strip() for c in vals}
                    header_row = row[0].row
                    break
            for row in rows:
                for c in row:
                    if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                        continue
                    v = c.value
                    if hasattr(v, "strftime"):
                        sval = v.strftime("%m/%d/%Y")
                    elif isinstance(v, float) and v.is_integer():
                        sval = str(int(v))
                    else:
                        sval = str(v)
                    label = header.get(c.column) if header_row and c.row != header_row else None
                    prefix = f"{label}: " if label else ""
                    buf.append(prefix)
                    pos += len(prefix)
                    self.cells.append((pos, pos + len(sval), c))
                    buf.append(sval)
                    pos += len(sval)
                    buf.append("\n")
                    pos += 1
            buf.append("\n")
            pos += 1
        self.text = "".join(buf)

    def analyze(self, options: AnalyzeOptions | None = None, analyzer: TextAnalyzer | None = None) -> list[Detection]:
        spans = (analyzer or get_analyzer()).analyze(self.text, options or AnalyzeOptions())
        return [_span_to_det(sp) for sp in spans if self._in_cell(sp.start, sp.end) or self._in_title(sp.start, sp.end)]

    def _in_title(self, a: int, b: int) -> bool:
        return any(s < b and a < e for s, e, _ in self.titles)

    def search(self, terms: list[SearchTerm]) -> list[Detection]:
        return [_span_to_det(sp) for sp in TextAnalyzer._user_rules(self.text, terms) if self._in_cell(sp.start, sp.end)]

    def _in_cell(self, a: int, b: int) -> bool:
        return any(s < b and a < e for s, e, _ in self.cells)

    def apply(self, detections: list[Detection], out_path: str | Path, style: str = "block") -> dict:
        n = 0
        edits: dict[int, list[tuple[int, int, Detection]]] = {}
        for d in detections:
            if not d.selected or d.start < 0:
                continue
            for idx, (s, e, cell) in enumerate(self.cells):
                if s < d.end and d.start < e:
                    edits.setdefault(idx, []).append((max(d.start, s) - s, min(d.end, e) - s, d))
        for idx, lst in edits.items():
            s, e, cell = self.cells[idx]
            val = self.text[s:e]
            for a, b, d in sorted(lst, key=lambda x: x[0], reverse=True):
                if d.partial:
                    for pa, pb in sorted(partial_ranges(val[a:b], d.partial), reverse=True):
                        val = val[:a + pa] + "█" * (pb - pa) + val[a + pb:]
                else:
                    val = val[:a] + redaction_text(d, style, val[a:b]) + val[b:]
                n += 1
            cell.value = val  # numbers become text: a redacted SSN must not survive as a number format
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.comment is not None:
                        c.comment = None
                        self.cleaned["comments"] = self.cleaned.get("comments", 0) + 1
                    if c.hyperlink is not None:
                        c.hyperlink = None
                        self.cleaned["hyperlinks"] = self.cleaned.get("hyperlinks", 0) + 1
            for hf in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter, ws.firstHeader, ws.firstFooter):
                for part in (hf.left, hf.center, hf.right):
                    part.text = None
        # sheet names can themselves be names ("Delgado payroll")
        tokens = {t.lower() for d in detections if d.selected for t in re.findall(r"[A-Za-z0-9]{3,}", d.text)}
        for k, (s, e, ws) in enumerate(self.titles):
            hit = any(d.selected and s < d.end and d.start < e for d in detections)
            if hit or tokens & {t.lower() for t in re.findall(r"[A-Za-z0-9]{3,}", ws.title)}:
                ws.title = f"Sheet{k + 1}"
                self.cleaned["sheet_titles"] = self.cleaned.get("sheet_titles", 0) + 1
        props = self.wb.properties
        for attr in ("creator", "lastModifiedBy", "title", "subject", "description", "keywords", "category"):
            setattr(props, attr, "" if attr in ("creator",) else None)
        self.wb.custom_doc_props.props.clear() if hasattr(self.wb, "custom_doc_props") else None
        out_path = Path(out_path)
        self.wb.save(str(out_path))
        report = {"redactions": n, "cleaned": self.cleaned}
        report["verification"] = verify_zip(out_path, [d for d in detections if d.selected and not d.partial])
        return report


def _norm(s: str) -> str:
    return re.sub(r"[^0-9a-z]", "", s.lower())


def verify_zip(path: Path, detections: list[Detection]) -> dict:
    """Search every XML part of the output package for redacted values (catches caches, rels, props)."""
    issues = []
    with zipfile.ZipFile(path) as z:
        blob = ""
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                xml = z.read(n).decode("utf-8", "ignore")
                blob += " " + re.sub(r"<[^>]+>", " ", xml)
        nb = _norm(blob)
        for d in detections:
            v = _norm(d.text)
            if len(v) >= 5 and (not v.isalpha() or len(v) >= 7) and v in nb:
                issues.append({"type": "value_still_present", "entity": d.entity})
        names = z.namelist()
        for n in names:
            if _DROP_PARTS.match(n):
                issues.append({"type": "sensitive_part_present", "part": n})
    return {"passed": not issues, "issues": issues}


def copy_file(src: Path, dst: Path) -> Path:
    shutil.copyfile(src, dst)
    return dst
