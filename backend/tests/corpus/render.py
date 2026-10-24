"""Render corpus documents to PDF (digital and simulated-scan variants) and DOCX."""
from __future__ import annotations

import io
import random
import textwrap
from pathlib import Path

import numpy as np
import pymupdf as fitz
from PIL import Image, ImageFilter

PAGE_W, PAGE_H = 612, 792  # US Letter
MARGIN = 54


def text_to_pdf(text: str, out: str | Path, mono: bool = False, fontsize: float = 9.5,
                metadata: dict | None = None) -> Path:
    doc = fitz.open()
    font = "cour" if mono else "helv"
    lh = fontsize * 1.32
    max_chars = int((PAGE_W - 2 * MARGIN) / (fontsize * (0.6 if mono else 0.5)))
    page = None
    y = 0.0

    def new_page():
        nonlocal page, y
        page = doc.new_page(width=PAGE_W, height=PAGE_H)
        y = MARGIN

    new_page()
    for block in text.split("\f"):
        if page is None or y > MARGIN:
            if block is not text.split("\f")[0]:
                new_page()
        for raw_line in block.split("\n"):
            wrapped = textwrap.wrap(raw_line, max_chars, replace_whitespace=False, drop_whitespace=False,
                                    break_long_words=False, break_on_hyphens=False) or [""]
            for ln in wrapped:
                if y + lh > PAGE_H - MARGIN:
                    new_page()
                if ln.strip():
                    page.insert_text(fitz.Point(MARGIN, y + fontsize), ln.replace("\t", "    "), fontsize=fontsize,
                                     fontname=font)
                y += lh
    doc.set_metadata(metadata or {"author": "Jane Q. Paralegal", "title": "Client file", "subject": "PI matter",
                                  "creator": "Word", "producer": "Obscura test corpus"})
    out = Path(out)
    doc.save(str(out))
    doc.close()
    return out


def scan_pdf(pdf_in: str | Path, out: str | Path, dpi: int = 200, seed: int = 0, quality: int = 70,
             max_rotate: float = 1.2, noise: float = 8.0, blur: float = 0.5) -> Path:
    """Simulate a scanner/fax: rasterise, rotate slightly, add noise and blur, JPEG-compress."""
    r = random.Random(seed)
    src = fitz.open(str(pdf_in))
    dst = fitz.open()
    for page in src:
        pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False, colorspace=fitz.csGRAY)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        angle = r.uniform(-max_rotate, max_rotate)
        img = img.rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=255)
        if blur:
            img = img.filter(ImageFilter.GaussianBlur(blur))
        arr = np.asarray(img).astype(np.float32)
        arr += np.random.default_rng(seed + page.number).normal(0, noise, arr.shape)
        img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=quality)
        p = dst.new_page(width=page.rect.width, height=page.rect.height)
        p.insert_image(p.rect, stream=buf.getvalue())
    out = Path(out)
    dst.save(str(out), garbage=3, deflate=True)
    return out


def text_to_docx(text: str, out: str | Path, title: str = "") -> Path:
    import docx
    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = title
    for block in text.split("\f"):
        for line in block.split("\n"):
            d.add_paragraph(line)
        d.add_page_break()
    d.core_properties.author = "Jane Q. Paralegal"
    d.core_properties.comments = "draft — contains client SSN"
    out = Path(out)
    d.save(str(out))
    return out
