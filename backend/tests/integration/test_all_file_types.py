"""Every supported file type, end to end through the real API: upload → analyze → export → download."""
import email.message
import io
import re
import subprocess
import time
import zipfile
from pathlib import Path

import docx
import openpyxl
import pymupdf as fitz
import pytest
from fastapi.testclient import TestClient

FIX = Path(__file__).resolve().parents[1] / "fixtures"
PW = "Str0ng!Passphrase"


@pytest.fixture(scope="module")
def api(tmp_path_factory):
    import importlib
    import os
    data = tmp_path_factory.mktemp("types")
    os.environ["OBSCURA_DATA"] = str(data)
    import obscura.config as cfg
    importlib.reload(cfg)
    import obscura.db as db
    import obscura.main as main
    import obscura.security as sec
    import obscura.storage as storage
    storage._STORAGE = storage._CIPHER = None
    storage.DATA_DIR = db.DATA_DIR = sec.DATA_DIR = main.DATA_DIR = cfg.DATA_DIR
    db._engine = db._Session = None
    sec._SECRET = None
    app = main.create_app()
    with TestClient(app) as c:
        pw = re.search(r"password: (\S+)", (cfg.DATA_DIR / "INITIAL_ADMIN_PASSWORD.txt").read_text()).group(1)
        t = c.post("/api/auth/login", json={"username": "admin", "password": pw}).json()["token"]
        t = c.post("/api/auth/change-password", json={"current_password": pw, "new_password": PW},
                   headers={"Authorization": f"Bearer {t}"}).json()["token"]
        h = {"Authorization": f"Bearer {t}"}
        cid = c.post("/api/cases", json={"name": "All types"}, headers=h).json()["id"]
        yield c, h, cid


def _job(c, h, jid, timeout=900):
    t = time.time()
    while time.time() - t < timeout:
        j = c.get(f"/api/jobs/{jid}", headers=h).json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.3)
    raise TimeoutError


def roundtrip(api, name, data, analyze=None, export=None):
    c, h, cid = api
    up = c.post(f"/api/cases/{cid}/files", files=[("files", (name, data, "application/octet-stream"))], headers=h).json()
    assert not up["errors"], up["errors"]
    fid = up["files"][0]["id"]
    j = _job(c, h, c.post(f"/api/files/{fid}/analyze", json=analyze or {"use_ml": False}, headers=h).json()["job_id"])
    assert j["status"] == "done", j
    dets = c.get(f"/api/files/{fid}/detections", headers=h).json()["detections"]
    j = _job(c, h, c.post(f"/api/files/{fid}/export", json=export or {}, headers=h).json()["job_id"])
    assert j["status"] == "done", j
    out = c.get(f"/api/exports/{j['result']['export_id']}/download", headers=h).content
    return fid, dets, j["result"], out, up["files"]


def test_pdf(api):
    d = fitz.open()
    d.new_page().insert_text((72, 72), "Patient: Maria Delgado  SSN 523-71-9921", fontsize=10)
    _, dets, res, out, _ = roundtrip(api, "a.pdf", d.tobytes())
    assert res["verified"] and b"523-71-9921" not in out
    assert "Delgado" not in fitz.open(stream=out, filetype="pdf")[0].get_text()


@pytest.mark.vision
def test_image_with_faces(api):
    _, dets, res, out, _ = roundtrip(api, "zidane.jpg", (FIX / "images" / "zidane.jpg").read_bytes(),
                                     analyze={"use_ml": False, "visual_classes": ["face", "head", "license_plate"]},
                                     export={"effect": "blur"})
    assert sum(1 for d in dets if d["entity"] == "FACE") >= 2 and res["verified"], res


@pytest.mark.vision
def test_video(api):
    c, h, cid = api
    fid, dets, res, out, _ = roundtrip(api, "clip.mp4", (FIX / "media" / "moving_faces.mp4").read_bytes(),
                                       analyze={"use_ml": False, "visual_classes": ["face"], "video_mode": "fast",
                                                "audio": True})
    tracks = c.get(f"/api/files/{fid}/video/tracks", headers=h).json()["tracks"]
    assert len(tracks) >= 2 and res["verified"], res
    tid = tracks[0]["id"]
    c.patch(f"/api/files/{fid}/video/tracks/{tid}", json={"trim_start": 10, "trim_end": 40}, headers=h)
    t = next(t for t in c.get(f"/api/files/{fid}/video/tracks", headers=h).json()["tracks"] if t["id"] == tid)
    assert t["first"] >= 10 and t["last"] <= 40
    assert len(out) > 10000


@pytest.mark.audio
def test_audio(api):
    c, h, cid = api
    fid, dets, res, out, files = roundtrip(api, "call.aiff", (FIX / "media" / "speech.aiff").read_bytes())
    assert files[0]["duration"] > 15
    assert {"SSN", "PHONE"} <= {d["entity"] for d in dets}
    assert res["verified"], res
    r = c.post(f"/api/files/{fid}/detections", json={"entity": "PERSON", "t0": 1.0, "t1": 2.5}, headers=h).json()
    assert r["text"]  # manual time-range redactions keep the spoken words


def test_email_with_attachment(api):
    att = fitz.open()
    att.new_page().insert_text((72, 72), "Records for Maria Delgado, DOB 01/15/1980", fontsize=10)
    m = email.message.EmailMessage()
    m["From"] = "Maria Delgado <maria.delgado@gmail.com>"
    m["To"] = "Jordan Brooks <jbrooks@brookslaw.com>"
    m["Subject"] = "My SSN is 523-71-9921"
    m.set_content("Call me at 512-555-0187.")
    m.add_attachment(att.tobytes(), maintype="application", subtype="pdf", filename="records.pdf")
    _, dets, res, out, files = roundtrip(api, "msg.eml", m.as_bytes())
    assert files[0]["kind"] == "pdf" and files[0]["info"]["converted_from"] == "email"
    assert any(f["parent_id"] == files[0]["id"] for f in files[1:])
    text = fitz.open(stream=out, filetype="pdf")[0].get_text()
    assert "523-71-9921" not in text and "512-555-0187" not in text and res["verified"]


def test_docx(api):
    d = docx.Document()
    d.add_paragraph("Client: Maria Delgado")
    d.add_paragraph("SSN: 523-71-9921")
    b = io.BytesIO()
    d.save(b)
    _, dets, res, out, _ = roundtrip(api, "a.docx", b.getvalue())
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        xml = z.read("word/document.xml").decode()
    assert "523-71-9921" not in xml and "Delgado" not in xml and res["verified"]


def test_xlsx(api):
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "SSN"])
    wb.active.append(["Maria Delgado", "523-71-9921"])
    b = io.BytesIO()
    wb.save(b)
    _, dets, res, out, _ = roundtrip(api, "a.xlsx", b.getvalue())
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        blob = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    assert "523-71-9921" not in blob and res["verified"]


@pytest.mark.parametrize("ext", ["txt", "csv", "json", "md"])
def test_text_formats(api, ext):
    body = {"txt": "Patient: Maria Delgado SSN 523-71-9921", "csv": "name,ssn\nMaria Delgado,523-71-9921",
            "json": '{"patient": "Maria Delgado", "ssn": "523-71-9921"}', "md": "# Note\nSSN: 523-71-9921"}[ext]
    _, dets, res, out, _ = roundtrip(api, f"a.{ext}", body.encode())
    assert b"523-71-9921" not in out and res["verified"]


def test_tiff_fax(api):
    from PIL import Image, ImageDraw, ImageFont
    pages = []
    for txt in ("FAX COVER", "Patient: Maria Delgado  SSN: 523-71-9921"):
        im = Image.new("L", (1700, 2200), 255)
        try:
            f = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 48)
        except OSError:
            f = ImageFont.load_default()
        ImageDraw.Draw(im).text((150, 300), txt, fill=0, font=f)
        pages.append(im)
    b = io.BytesIO()
    pages[0].save(b, "TIFF", save_all=True, append_images=pages[1:], dpi=(200, 200))
    _, dets, res, out, files = roundtrip(api, "fax.tiff", b.getvalue())
    assert files[0]["kind"] == "pdf" and files[0]["page_count"] == 2
    assert any(d["entity"] == "SSN" for d in dets) and res["verified"]


def test_failed_analysis_marks_file(api):
    c, h, cid = api
    up = c.post(f"/api/cases/{cid}/files", files=[("files", ("broken.pdf", b"%PDF-1.7 not really", "application/pdf"))],
                headers=h).json()
    if up["errors"]:
        return  # rejected at upload: also fine
    fid = up["files"][0]["id"]
    _job(c, h, c.post(f"/api/files/{fid}/analyze", json={"use_ml": False}, headers=h).json()["job_id"])
    f = c.get(f"/api/files/{fid}", headers=h).json()
    assert f["status"] == "error" and f["error"]
