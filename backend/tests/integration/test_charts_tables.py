"""Written-record redaction end to end: medical charts, flowsheets, MARs, lab and billing tables, chart graphics
(line / bar / pie images, vector charts with rotated axis labels), scanned charts, and the same content in
Word, PowerPoint and Excel.

Every case builds a realistic file, runs detection, applies the redaction, then inspects the *output file*
(text extraction, OCR of the rendered pages and pictures, raw package XML) to prove each identifier is gone
while the clinical content a reader needs is still there."""
from __future__ import annotations

import io
import re
import zipfile

import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw, ImageFont

from obscura.detection.text.engine import AnalyzeOptions
from obscura.redact.model import ExportOptions
from obscura.redact.pdf import PdfDocument

RULES = AnalyzeOptions(use_spacy=False, use_gliner=False)
FULL = AnalyzeOptions()  # the product default: rules + local spaCy + GLiNER
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"


def _norm(s: str) -> str:
    return re.sub(r"[^0-9a-z]", "", s.lower())


def _ocr(img: Image.Image) -> str:
    """Read the output the way an investigator would: two different OCR passes, combined (the product's
    own table-aware OCR, plus plain Tesseract)."""
    import pytesseract
    from obscura.redact import ocr
    return ocr.words_to_text(ocr.ocr_image(img, dpi=300))[0] + "\n" + pytesseract.image_to_string(img, config="--psm 6")


# =====================================================================================================
# 1. Medical charts / record tables as native PDF text
# =====================================================================================================
# Each chart: header block + a table drawn in columns. PII must disappear; the listed clinical anchors stay.

CHARTS = {
    "vitals_flowsheet": {
        "lines": [
            "MERCY RIDGE MEDICAL CENTER - VITAL SIGNS FLOWSHEET",
            "Patient: OKONKWO, CHIDINMA A        MRN: 00731945        DOB: 09/14/1978",
            "Room 512-B   Attending: Dr. Harold Lindqvist   Encounter #: E44871290",
            "",
            "Time     Temp    Pulse   Resp   BP         SpO2   Pain   Initials",
            "0600     98.6    88      18     132/84     97%    6/10   RT",
            "1000     99.1    92      20     128/80     96%    5/10   RT",
            "1400     98.9    84      16     124/78     98%    4/10   KMB",
            "1800     98.4    80      16     122/76     98%    3/10   KMB",
            "",
            "Nurse: Rachel Tomczak, RN   Phone: (614) 293-5571",
        ],
        "pii": ["OKONKWO, CHIDINMA A", "00731945", "09/14/1978", "Harold Lindqvist", "E44871290",
                "Rachel Tomczak", "(614) 293-5571"],
        "keep": ["Pulse", "SpO2", "132/84", "Temp"],
    },
    "medication_administration_record": {
        "lines": [
            "MEDICATION ADMINISTRATION RECORD (MAR)",
            "Name: Delacroix, Beatrice M.     MR#: 55120874     DOB: 02/28/1951     Allergies: PCN",
            "Prescriber: Anand Krishnaswamy, MD   NPI 1730245569",
            "",
            "Medication                 Dose     Route   0800      1200      2000",
            "Lisinopril                 10 mg    PO      given     -         given",
            "Metformin                  500 mg   PO      given     given     given",
            "Acetaminophen              650 mg   PO      held      given     -",
            "",
            "Administered by: Jolene Swiftwater, LPN (JS)   Verified by: Tomas Ibarra-Vega, RN",
        ],
        # (a prescriber's NPI is a public provider number, not patient PHI, so it is not required here)
        "pii": ["Delacroix, Beatrice M", "55120874", "02/28/1951", "Anand Krishnaswamy",
                "Jolene Swiftwater", "Tomas Ibarra-Vega"],
        "keep": ["Lisinopril", "Route", "650 mg"],
    },
    "lab_results": {
        "lines": [
            "LAKESIDE REGIONAL LABORATORY - FINAL REPORT",
            "Patient Name: Haddad, Yusuf K       Patient ID: LRL-0094412      Sex: M",
            "Date of Birth: 11/03/1966           Collected: 05/12/2025 07:40",
            "Ordering Provider: Dr. Meredith Oyelaran     Account #: 7740022918",
            "",
            "Test                    Result     Flag    Reference",
            "Hemoglobin A1c          8.2 %      H       4.0 - 5.6",
            "Glucose, fasting        162 mg/dL  H       70 - 99",
            "Creatinine              1.1 mg/dL          0.7 - 1.3",
            "",
            "Report to: yusuf.haddad66@outlook.com",
        ],
        "pii": ["Haddad, Yusuf K", "LRL-0094412", "11/03/1966", "Meredith Oyelaran", "7740022918",
                "yusuf.haddad66@outlook.com"],
        "keep": ["Hemoglobin A1c", "Reference", "Creatinine"],
    },
    "problem_list": {
        "lines": [
            "PROBLEM LIST / ALLERGIES",
            "PATIENT: SANTIAGO-REYES, LUZ MARINA     MRN 00918334     DOB 07/22/1990",
            "Home: 2215 Calloway Ridge Dr, Marietta, GA 30064     Cell 770-555-0193",
            "",
            "#   Problem                          Onset      Status",
            "1   Lumbar radiculopathy             03/2025    Active",
            "2   Post-traumatic headache          03/2025    Active",
            "3   Anxiety                          2019       Active",
            "",
            "Emergency contact: Ernesto Reyes (husband) 770-555-0148",
        ],
        "pii": ["SANTIAGO-REYES, LUZ MARINA", "00918334", "07/22/1990", "2215 Calloway Ridge Dr", "770-555-0193",
                "Ernesto Reyes", "770-555-0148"],
        "keep": ["Lumbar radiculopathy", "Onset", "Status"],
    },
    "nursing_notes_timeline": {
        "lines": [
            "NURSING NARRATIVE NOTES",
            "Pt: Wojcik, Dariusz    MRN: 66231078    DOB: 12/01/1983    Unit: 4 North",
            "",
            "05/14 0215  Pt resting. Wife Agnieszka at bedside, asked to call 312-555-0177 if changes.",
            "05/14 0630  Pt reports pain 7/10. Dr. Feldstein notified. Morphine 2 mg IV given.",
            "05/14 0900  Pt states he works at Lakeshore Steel Fabrication as a welder.",
            "05/14 1130  Daughter Zofia (age 9) visited with grandmother.",
            "",
            "Signed: Kimberly Ashworth-Nnamdi, RN",
        ],
        "pii": ["Wojcik, Dariusz", "66231078", "12/01/1983", "Agnieszka", "312-555-0177", "Feldstein",
                "Lakeshore Steel Fabrication", "Zofia", "Kimberly Ashworth-Nnamdi"],
        "keep": ["Pt resting", "pain 7/10"],
    },
    "itemized_bill": {
        "lines": [
            "ST. ANSELM HOSPITAL - ITEMIZED STATEMENT",
            "Guarantor: Brennan, Colm P      Guarantor Acct: 0048821906      Statement Date: 06/01/2025",
            "Patient Account #: 9920014477   Insurance: Buckeye Health Plan  Member ID: BHP77310429",
            "Group #: 448210",
            "",
            "Date     Rev   CPT     Description                    Charges",
            "03/14    0450  99285   ED visit level 5               2,310.00",
            "03/14    0350  70450   CT head w/o contrast           1,845.00",
            "03/14    0320  72125   CT cervical spine              1,960.00",
            "",
            "Mail payment to: Colm Brennan, 418 Keller Ave, Dayton, OH 45410",
        ],
        "pii": ["Brennan, Colm P", "0048821906", "9920014477", "BHP77310429", "448210", "Colm Brennan",
                "418 Keller Ave"],
        "keep": ["CT head w/o contrast", "Charges", "99285"],
    },
    "pediatric_growth_header": {
        "lines": [
            "PEDIATRIC GROWTH CHART - BOYS 2 TO 20 YEARS",
            "Name: Mateo Alvarez-Quintero (minor)    DOB: 08/19/2016    MRN: 30087412",
            "Parent/Guardian: Lucia Quintero    Phone: (505) 555-0166",
            "School: Monte Vista Elementary, 3rd grade",
            "",
            "Age      Height (in)   Weight (lb)   BMI",
            "6 y      45.5          46            15.6",
            "7 y      48.0          52            15.9",
            "8 y      50.2          60            16.7",
        ],
        "pii": ["Mateo Alvarez-Quintero", "08/19/2016", "30087412", "Lucia Quintero", "(505) 555-0166",
                "Monte Vista Elementary"],
        "keep": ["Height (in)", "BMI", "45.5"],
    },
    "pharmacy_profile": {
        "lines": [
            "CORNERSTONE PHARMACY #0412 - PATIENT PROFILE",
            "PATIENT: NAKAMURA, KENJI      DOB: 04/05/1972      PHONE: 206-555-0134",
            "ADDRESS: 7712 Linden Ave N, Seattle, WA 98103",
            "INSURANCE: RX GRP 9KX221   MEMBER ID: ZXW448810331",
            "",
            "Fill Date   Rx #       Drug                       Qty   Prescriber",
            "04/02/25    6612049    Gabapentin 300 mg cap      90    T. Abernethy",
            "04/02/25    6612050    Cyclobenzaprine 10 mg      30    T. Abernethy",
            "05/01/25    6615231    Gabapentin 300 mg cap      90    T. Abernethy",
        ],
        "pii": ["NAKAMURA, KENJI", "04/05/1972", "206-555-0134", "7712 Linden Ave N", "ZXW448810331"],
        "keep": ["Fill Date", "Qty", "Prescriber"],
    },
    "insurance_eligibility": {
        "lines": [
            "ELIGIBILITY & BENEFITS VERIFICATION",
            "Subscriber: Thibodeaux, Renee L     Subscriber ID: LAH992017734     DOB: 01/30/1985",
            "Dependent: Thibodeaux, Jacques (son)   Dependent DOB: 06/11/2012",
            "Policy #: 77-4410923-01     Claim #: CLM-2025-0081472",
            "",
            "Benefit                 In-network      Out-of-network",
            "Deductible              $1,500          $3,000",
            "Coinsurance             20%             40%",
        ],
        "pii": ["Thibodeaux, Renee L", "LAH992017734", "01/30/1985", "Thibodeaux, Jacques", "06/11/2012",
                "77-4410923-01", "CLM-2025-0081472"],
        "keep": ["Deductible", "Coinsurance", "In-network"],
    },
    "pain_diary": {
        "lines": [
            "DAILY PAIN DIARY - completed by patient",
            "Name: Priya Raghunathan     Claim No. WC-55-120934     Employer: Northgate Logistics",
            "",
            "Date    Pain   Activities / notes",
            "6/02    7      Could not lift my daughter Anika. Skipped yoga at the YMCA on Delmar.",
            "6/03    6      Drove to PT at Midtown Rehab. Neighbor Mrs. Okafor helped with groceries.",
            "6/04    8      Woke at 3am. Called Dr. Whitcombe's office 314-555-0102.",
        ],
        "pii": ["Priya Raghunathan", "WC-55-120934", "Northgate Logistics", "Anika", "Okafor", "Whitcombe",
                "314-555-0102"],
        "keep": ["Activities / notes", "Pain"],
    },
}


def _chart_pdf(path, lines: list[str], fontsize: float = 8.5) -> None:
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 60
    for ln in lines:
        if ln:
            page.insert_text((40, y), ln, fontsize=fontsize, fontname="cour")
        y += fontsize * 1.6
    # table rules like a real flowsheet
    page.draw_line((40, 50), (572, 50))
    page.draw_line((40, y + 4), (572, y + 4))
    doc.save(str(path))


@pytest.fixture(scope="module")
def redacted_charts(tmp_path_factory):
    d = tmp_path_factory.mktemp("charts")
    out = {}
    for name, spec in CHARTS.items():
        src, dst = d / f"{name}.pdf", d / f"{name}_red.pdf"
        _chart_pdf(src, spec["lines"])
        doc = PdfDocument(src)
        dets = doc.analyze(RULES)
        rep = doc.apply(dets, dst, ExportOptions())
        text = " ".join(p.get_text() for p in fitz.open(str(dst)))
        out[name] = {"text": text, "report": rep}
    return out


PII_CASES = [(n, v) for n, s in CHARTS.items() for v in s["pii"]]
KEEP_CASES = [(n, v) for n, s in CHARTS.items() for v in s["keep"]]


@pytest.mark.parametrize("chart,value", PII_CASES, ids=[f"{c}:{v[:18]}" for c, v in PII_CASES])
def test_chart_identifier_removed(redacted_charts, chart, value):
    assert _norm(value) not in _norm(redacted_charts[chart]["text"])


@pytest.mark.parametrize("chart,value", KEEP_CASES, ids=[f"{c}:{v[:18]}" for c, v in KEEP_CASES])
def test_chart_clinical_content_kept(redacted_charts, chart, value):
    assert _norm(value) in _norm(redacted_charts[chart]["text"])


@pytest.mark.parametrize("chart", list(CHARTS))
def test_chart_export_verified(redacted_charts, chart):
    assert redacted_charts[chart]["report"]["verification"]["passed"], redacted_charts[chart]["report"]


# =====================================================================================================
# 2. Chart graphics: images inside an otherwise digital PDF page
# =====================================================================================================

def _plot(kind: str) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(6.4, 3.4), dpi=170)
    if kind == "line":
        ax.plot([1, 2, 3, 4, 5, 6], [8, 7, 7, 6, 5, 4], marker="o", label="Patient: Chidinma Okonkwo")
        ax.set_title("Pain scores - MRN 00731945 - DOB 09/14/1978", fontsize=11)
        ax.set_xlabel("Visit")
        ax.legend()
    elif kind == "bar":
        ax.bar(["Mar", "Apr", "May"], [2310, 1845, 1960])
        ax.set_title("Medical expenses - Colm Brennan - Acct 0048821906", fontsize=11)
    elif kind == "pie":
        ax.pie([45, 30, 25], labels=["Lost wages", "Medical", "Future care"])
        ax.set_title("Damages summary for Renee Thibodeaux (SSN 438-22-6719)", fontsize=10)
    elif kind == "table":
        ax.axis("off")
        tb = ax.table(cellText=[["Yusuf Haddad", "11/03/1966", "LRL-0094412"], ["Glucose", "162", "H"]],
                      colLabels=["Patient", "DOB", "Patient ID"], bbox=[0, 0.1, 1, 0.8])
        tb.auto_set_font_size(False)
        tb.set_fontsize(13)
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


IMAGE_CHARTS = {
    "line": (["Okonkwo", "00731945", "09/14/1978"], ["Pain scores", "Patient"]),
    "bar": (["Brennan", "0048821906"], ["Medical expenses"]),
    "pie": (["Thibodeaux", "438-22-6719"], ["Lost wages", "Damages summary"]),
    "table": (["Haddad", "11/03/1966", "LRL-0094412"], ["Glucose"]),
}


@pytest.fixture(scope="module")
def redacted_image_charts(tmp_path_factory):
    d = tmp_path_factory.mktemp("imgcharts")
    out = {}
    for kind in IMAGE_CHARTS:
        src, dst = d / f"{kind}.pdf", d / f"{kind}_red.pdf"
        doc = fitz.open()
        page = doc.new_page(width=612, height=792)
        page.insert_text((50, 60), "Exhibit 4 - summary chart prepared for mediation.", fontsize=11)
        page.insert_image(fitz.Rect(50, 90, 50 + 420, 90 + 223), stream=_plot(kind))  # ~19% of the page
        doc.save(str(src))
        pd = PdfDocument(src)
        dets = pd.analyze(AnalyzeOptions())
        rep = pd.apply(dets, dst, ExportOptions())
        pix = fitz.open(str(dst))[0].get_pixmap(dpi=300)
        out[kind] = {"ocr": _ocr(Image.open(io.BytesIO(pix.tobytes("png")))), "report": rep,
                     "found": {d.text for d in dets if d.selected}}
    return out


IMG_PII = [(k, v) for k, (pii, _) in IMAGE_CHARTS.items() for v in pii]
IMG_KEEP = [(k, v) for k, (_, keep) in IMAGE_CHARTS.items() for v in keep]


@pytest.mark.parametrize("kind,value", IMG_PII, ids=[f"{k}:{v}" for k, v in IMG_PII])
def test_chart_image_identifier_burned_out(redacted_image_charts, kind, value):
    assert _norm(value) not in _norm(redacted_image_charts[kind]["ocr"])


@pytest.mark.parametrize("kind,value", IMG_KEEP, ids=[f"{k}:{v}" for k, v in IMG_KEEP])
def test_chart_image_rest_still_readable(redacted_image_charts, kind, value):
    assert _norm(value) in _norm(redacted_image_charts[kind]["ocr"])


# =====================================================================================================
# 3. Vector charts: labels are real PDF text, some rotated
# =====================================================================================================

@pytest.fixture(scope="module")
def vector_chart(tmp_path_factory):
    d = tmp_path_factory.mktemp("vec")
    src, dst = d / "vec.pdf", d / "vec_red.pdf"
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((60, 60), "Range of motion - Patient: Kowalczyk, Darek  MRN: 44120938", fontsize=11)
    page.draw_rect(fitz.Rect(100, 100, 500, 400))
    for i, h in enumerate([120, 180, 220, 260]):
        page.draw_rect(fitz.Rect(130 + i * 90, 400 - h, 180 + i * 90, 400), fill=(0.3, 0.5, 0.8))
    page.insert_text((80, 330), "Degrees - Darek Kowalczyk", fontsize=10, rotate=90)  # y-axis label
    page.insert_text((200, 425), "Visits with Dr. Ingrid Solberg (fax 763-555-0181)", fontsize=10)
    page.insert_text((360, 120), "SSN 529-48-1173", fontsize=8)  # data label someone typed on the chart
    doc.save(str(src))
    pd = PdfDocument(src)
    rep = pd.apply(pd.analyze(RULES), dst, ExportOptions())
    return " ".join(p.get_text() for p in fitz.open(str(dst))), rep


@pytest.mark.parametrize("value", ["Kowalczyk, Darek", "44120938", "Darek Kowalczyk", "Ingrid Solberg",
                                   "763-555-0181", "529-48-1173"])
def test_vector_chart_labels_removed(vector_chart, value):
    assert _norm(value) not in _norm(vector_chart[0])


def test_vector_chart_axis_words_kept(vector_chart):
    assert "degrees" in vector_chart[0].lower() and "range of motion" in vector_chart[0].lower()


# =====================================================================================================
# 4. Scanned charts (the whole page is a picture)
# =====================================================================================================

SCAN_SPECS = ["vitals_flowsheet", "lab_results", "pharmacy_profile"]


@pytest.fixture(scope="module")
def scanned_charts(tmp_path_factory):
    d = tmp_path_factory.mktemp("scan")
    font = ImageFont.truetype(FONT, 26)
    out = {}
    for name in SCAN_SPECS:
        img = Image.new("RGB", (2100, 1500), "white")
        dr = ImageDraw.Draw(img)
        for i, ln in enumerate(CHARTS[name]["lines"]):
            dr.text((80, 80 + i * 44), ln, fill="black", font=font)
        buf = io.BytesIO()
        img.save(buf, "PNG")
        src, dst = d / f"{name}.pdf", d / f"{name}_red.pdf"
        doc = fitz.open()
        page = doc.new_page(width=792, height=566)
        page.insert_image(page.rect, stream=buf.getvalue())
        doc.save(str(src))
        pd = PdfDocument(src)
        rep = pd.apply(pd.analyze(RULES), dst, ExportOptions())
        pix = fitz.open(str(dst))[0].get_pixmap(dpi=250)
        out[name] = (_ocr(Image.open(io.BytesIO(pix.tobytes("png")))), rep)
    return out


SCAN_CASES = [(n, v) for n in SCAN_SPECS for v in CHARTS[n]["pii"] if re.search(r"\d{4}", v)]


@pytest.mark.parametrize("chart,value", SCAN_CASES, ids=[f"{c}:{v}" for c, v in SCAN_CASES])
def test_scanned_chart_numbers_removed(scanned_charts, chart, value):
    assert re.sub(r"\D", "", value) not in re.sub(r"\D", "", scanned_charts[chart][0])


@pytest.mark.parametrize("chart", SCAN_SPECS)
def test_scanned_chart_keeps_table_headers(scanned_charts, chart):
    assert any(_norm(k) in _norm(scanned_charts[chart][0]) for k in CHARTS[chart]["keep"])


# =====================================================================================================
# 5. Word, PowerPoint and Excel: tables, chart data and pictures
# =====================================================================================================

def _package_text(path) -> str:
    z = zipfile.ZipFile(path)
    return " ".join(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8", "ignore")) for n in z.namelist()
                    if n.endswith((".xml", ".rels")))


@pytest.fixture(scope="module")
def docx_table(tmp_path_factory):
    import docx
    from obscura.redact.office import OoxmlDocument
    d = tmp_path_factory.mktemp("docx")
    doc = docx.Document()
    doc.add_paragraph("Treatment summary prepared for counsel")
    t = doc.add_table(rows=1, cols=4)
    for c, h in zip(t.rows[0].cells, ["Patient", "DOB", "MRN", "Provider"]):
        c.text = h
    for row in [("Okonkwo, Chidinma", "09/14/1978", "00731945", "Dr. Harold Lindqvist"),
                ("Haddad, Yusuf", "11/03/1966", "LRL-0094412", "Dr. Meredith Oyelaran")]:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    doc.add_paragraph("Phone on file: (614) 293-5571. SSN: 523-71-9921.")
    src, dst = d / "t.docx", d / "t_red.docx"
    doc.save(str(src))
    od = OoxmlDocument(src)
    rep = od.apply(od.analyze(AnalyzeOptions()), dst)
    return _package_text(dst), rep


@pytest.mark.parametrize("value", ["Okonkwo, Chidinma", "09/14/1978", "00731945", "Harold Lindqvist", "Haddad, Yusuf",
                                   "11/03/1966", "LRL-0094412", "Meredith Oyelaran", "(614) 293-5571", "523-71-9921"])
def test_docx_table_values_removed(docx_table, value):
    assert _norm(value) not in _norm(docx_table[0])


def test_docx_table_headers_kept(docx_table):
    assert all(h in docx_table[0] for h in ("Patient", "DOB", "MRN", "Provider"))


@pytest.fixture(scope="module")
def pptx_chart(tmp_path_factory):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from obscura.redact.office import OoxmlDocument
    d = tmp_path_factory.mktemp("pptx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Recovery of Renee Thibodeaux (DOB 01/30/1985)"
    cd = CategoryChartData()
    cd.categories = ["Visit 1 - 03/14/2025", "Visit 2 - 04/02/2025", "Visit 3 - 05/01/2025"]
    cd.add_series("Renee Thibodeaux pain score", (8, 6, 4))
    cd.add_series("Member ID LAH992017734", (1, 2, 3))
    s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
    s.notes_slide.notes_text_frame.text = "Speaker notes: call Renee at 504-555-0139 before mediation."
    src, dst = d / "c.pptx", d / "c_red.pptx"
    prs.save(str(src))
    od = OoxmlDocument(src)
    rep = od.apply(od.analyze(AnalyzeOptions()), dst)
    return _package_text(dst), zipfile.ZipFile(dst).namelist(), rep


@pytest.mark.parametrize("value", ["Thibodeaux", "01/30/1985", "LAH992017734", "504-555-0139"])
def test_pptx_chart_and_notes_values_removed(pptx_chart, value):
    assert _norm(value) not in _norm(pptx_chart[0])


def test_pptx_chart_embedded_workbook_dropped(pptx_chart):
    # the chart's data workbook is a full copy of the series names; it must not ship
    assert not [n for n in pptx_chart[1] if n.startswith("ppt/embeddings/")]


def test_pptx_chart_keeps_scores(pptx_chart):
    assert "pain score" in pptx_chart[0].lower()


@pytest.fixture(scope="module")
def xlsx_sheet(tmp_path_factory):
    import openpyxl
    from obscura.redact.office import XlsxDocument
    d = tmp_path_factory.mktemp("xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medical specials"
    ws.append(["Patient", "DOB", "SSN", "Provider", "Date of service", "Charges"])
    ws.append(["Brennan, Colm", "05/02/1979", "438-22-6719", "Dr. Ingrid Solberg", "03/14/2025", 2310])
    ws.append(["Brennan, Colm", "05/02/1979", "438-22-6719", "Midtown Rehab", "04/02/2025", 450])
    src, dst = d / "s.xlsx", d / "s_red.xlsx"
    wb.save(str(src))
    xd = XlsxDocument(src)
    rep = xd.apply(xd.analyze(RULES), dst)
    vals = " ".join(str(c.value) for row in openpyxl.load_workbook(dst).active.iter_rows() for c in row if c.value)
    return vals + " " + _package_text(dst), rep


@pytest.mark.parametrize("value", ["Brennan", "05/02/1979", "438-22-6719", "Ingrid Solberg"])
def test_xlsx_values_removed(xlsx_sheet, value):
    assert _norm(value) not in _norm(xlsx_sheet[0])


def test_xlsx_numbers_kept(xlsx_sheet):
    assert "2310" in xlsx_sheet[0] and "Charges" in xlsx_sheet[0]


@pytest.fixture(scope="module")
def docx_screenshot(tmp_path_factory):
    import docx
    from docx.shared import Inches
    from obscura.redact.office import OoxmlDocument
    d = tmp_path_factory.mktemp("shot")
    img = Image.new("RGB", (1000, 330), "white")
    dr = ImageDraw.Draw(img)
    f = ImageFont.truetype(FONT, 36)
    dr.text((24, 24), "Patient: Luz Santiago-Reyes", fill="black", font=f)
    dr.text((24, 100), "Call me at 770-555-0193", fill="black", font=f)
    dr.text((24, 176), "my SSN is 523-71-9921", fill="black", font=f)
    dr.text((24, 252), "See you at therapy Tuesday", fill="black", font=f)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = docx.Document()
    doc.add_paragraph("Exhibit B: screenshot of text messages from SANTIAGO-REYES, LUZ (client).")
    doc.add_picture(io.BytesIO(buf.getvalue()), width=Inches(5.5))
    src, dst = d / "s.docx", d / "s_red.docx"
    doc.save(str(src))
    od = OoxmlDocument(src)
    rep = od.apply(od.analyze(AnalyzeOptions()), dst)
    z = zipfile.ZipFile(dst)
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    return _ocr(Image.open(io.BytesIO(z.read(media[0])))), rep


@pytest.mark.parametrize("value", ["770-555-0193", "523-71-9921", "Santiago"])
def test_docx_picture_text_burned_out(docx_screenshot, value):
    assert _norm(value) not in _norm(docx_screenshot[0])


def test_docx_picture_rest_readable_and_verified(docx_screenshot):
    assert "therapy" in docx_screenshot[0].lower()
    assert docx_screenshot[1]["verification"]["passed"], docx_screenshot[1]
