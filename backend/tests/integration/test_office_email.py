"""Word / PowerPoint / Excel / email redaction: every carrier in the package must be clean."""
import email.message
import re
import zipfile

import docx
import openpyxl
import pymupdf as fitz
import pytest
from docx.oxml.ns import qn
from lxml import etree

from obscura.detection.text.engine import AnalyzeOptions
from obscura.redact.email_doc import email_to_pdf, parse_eml
from obscura.redact.office import OoxmlDocument, XlsxDocument
from obscura.redact.pdf import PdfDocument

RULES = AnalyzeOptions(use_spacy=False, use_gliner=False)
SSN, NAME, PHONE, EMAIL = "523-71-9921", "Maria Delgado", "512-555-0187", "maria.delgado@gmail.com"


def package_text(path) -> str:
    out = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                out.append(z.read(n).decode("utf-8", "ignore"))
    return "\n".join(out)


def _w(tag):
    return qn(f"w:{tag}")


@pytest.fixture
def tricky_docx(tmp_path):
    d = docx.Document()
    d.add_paragraph(f"Patient: {NAME}")
    p = d.add_paragraph("SSN: ")
    p.add_run(SSN[:4])  # value split across runs
    p.add_run(SSN[4:])
    d.add_paragraph("Clinical summary: lumbar strain, improving.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Phone", "Email"
    t.cell(1, 0).text, t.cell(1, 1).text = PHONE, EMAIL
    sec = d.sections[0]
    sec.header.paragraphs[0].text = f"Confidential — {NAME}"
    sec.footer.paragraphs[0].text = f"MRN: 00483921"
    # tracked deletion (invisible but in the file)
    para = d.add_paragraph("Deleted text follows: ")
    dl = etree.SubElement(para._p, _w("del"))
    r = etree.SubElement(dl, _w("r"))
    dt = etree.SubElement(r, _w("delText"))
    dt.text = "SSN 412-55-9023 removed by edit"
    # hidden run
    hp = d.add_paragraph("Visible. ")
    run = hp.add_run("Hidden note: client cell 713-555-0144")
    run.font.hidden = True
    # VML text box
    tb = d.add_paragraph()
    pict = etree.fromstring(
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textbox><w:txbxContent><w:p><w:r>'
        '<w:t>Driver\'s License #: 48392011</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')
    r = etree.SubElement(tb._p, _w("r"))
    r.append(pict)
    # field code with mailto
    fp = d.add_paragraph()
    r1 = etree.SubElement(fp._p, _w("r"))
    it = etree.SubElement(r1, _w("instrText"))
    it.text = f' HYPERLINK "mailto:{EMAIL}" '
    d.core_properties.author = NAME
    d.core_properties.comments = f"contains SSN {SSN}"
    d.core_properties.last_modified_by = "Jane Paralegal"
    path = tmp_path / "tricky.docx"
    d.save(str(path))
    return path


def test_docx_all_carriers(tricky_docx, tmp_path):
    doc = OoxmlDocument(tricky_docx)
    assert "412-55-9023" not in doc.text and "713-555-0144" not in doc.text  # removed before analysis
    dets = doc.analyze(RULES)
    out = tmp_path / "o.docx"
    rep = doc.apply(dets, out)
    blob = package_text(out)
    for v in (SSN, "412-55-9023", "713-555-0144", PHONE, EMAIL, "48392011", "00483921", "Delgado", "Jane Paralegal"):
        assert v not in blob, v
    assert "lumbar strain" in blob and rep["verification"]["passed"]
    assert rep["cleaned"].get("tracked_deletions") and rep["cleaned"].get("hidden_runs") and rep["cleaned"].get("field_codes")
    docx.Document(str(out))  # still a valid Word file


@pytest.mark.parametrize("style,marker", [("block", "█"), ("entity", "[SOCIAL SECURITY NUMBER]"), ("exemption", "[(b)(6)]")])
def test_docx_styles(tricky_docx, tmp_path, style, marker):
    doc = OoxmlDocument(tricky_docx)
    out = tmp_path / f"o_{style}.docx"
    doc.apply(doc.analyze(RULES), out, style=style)
    assert marker in package_text(out)


def test_docx_partial(tmp_path):
    d = docx.Document()
    d.add_paragraph(f"SSN: {SSN}")
    src = tmp_path / "p.docx"
    d.save(str(src))
    doc = OoxmlDocument(src)
    dets = doc.analyze(AnalyzeOptions(profile="frcp_5_2", use_spacy=False, use_gliner=False))
    out = tmp_path / "o.docx"
    doc.apply(dets, out)
    t = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert t.endswith("9921") and "523" not in t


def test_pptx(tmp_path):
    from pptx import Presentation
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = f"Claimant: {NAME}"
    s.placeholders[1].text = f"Claimant SSN: {SSN}\nPhone: {PHONE}"
    s.notes_slide.notes_text_frame.text = f"Speaker notes: DOB 01/15/1980, email {EMAIL}"
    pr.core_properties.author = NAME
    src = tmp_path / "deck.pptx"
    pr.save(str(src))
    doc = OoxmlDocument(src)
    out = tmp_path / "o.pptx"
    rep = doc.apply(doc.analyze(RULES), out)
    blob = package_text(out)
    for v in (SSN, PHONE, EMAIL, "01/15/1980", "Delgado"):
        assert v not in blob, v
    assert rep["verification"]["passed"]


def test_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delgado payroll"
    ws.append(["Employee Name", "SSN", "DOB", "Rate", "Routing", "Account", "Notes"])
    ws.append([NAME, "523-71-9921", "04/17/1986", 22.5, "021000021", 483920117, "ok"])
    ws.append(["Kevin O'Brien", 412559023, "1979-02-02", 31.0, "026009593", 7700112233, f"call {PHONE}"])
    ws["A2"].comment = openpyxl.comments.Comment(f"her email {EMAIL}", "HR")
    hidden = wb.create_sheet("hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["SSN", "601-22-4471"])
    wb.properties.creator = "Jane HR"
    src = tmp_path / "pay.xlsx"
    wb.save(str(src))
    x = XlsxDocument(src)
    out = tmp_path / "o.xlsx"
    rep = x.apply(x.analyze(RULES), out)
    blob = package_text(out)
    for v in ("523-71-9921", "412559023", "601-22-4471", "483920117", "7700112233", PHONE, EMAIL, "Delgado", "Jane HR"):
        assert v not in blob, v
    w2 = openpyxl.load_workbook(str(out))
    assert w2.worksheets[0].cell(2, 4).value == 22.5  # unflagged numbers stay numbers
    assert "Delgado" not in w2.sheetnames[0]
    assert rep["verification"]["passed"]


def _eml(tmp_path, attach_pdf: bytes | None = None):
    m = email.message.EmailMessage()
    m["From"] = f"{NAME} <{EMAIL}>"
    m["To"] = "Jordan Brooks <jbrooks@brookslaw.com>"
    m["Cc"] = "Priya Venkataraman <pv@brookslaw.com>"
    m["Subject"] = f"My claim — SSN {SSN}"
    m["Date"] = "Mon, 20 Jun 2022 10:15:00 -0500"
    m["Received"] = "from mail.example (203.0.113.77) by mx.brookslaw.com"
    m["X-Originating-IP"] = "[203.0.113.77]"
    m.set_content(f"Hi Jordan,\n\nMy new number is {PHONE}. I live at 4417 Oakridge Dr, Austin, TX 78745.\n\nThanks,\nMaria")
    m.add_alternative(f"<p>Hi Jordan,</p><p>My new number is <b>{PHONE}</b>.</p>", subtype="html")
    if attach_pdf:
        m.add_attachment(attach_pdf, maintype="application", subtype="pdf", filename="records.pdf")
    p = tmp_path / "msg.eml"
    p.write_bytes(m.as_bytes())
    return p


def test_email_print_and_redaction(tmp_path):
    att = fitz.open()
    att.new_page().insert_text((72, 72), f"Patient: {NAME}  SSN {SSN}", fontsize=11)
    parsed = parse_eml(_eml(tmp_path, att.tobytes()).read_bytes())
    assert parsed.attachments and parsed.attachments[0][0] == "records.pdf"
    hdrs = dict(parsed.headers)
    assert "Received" not in hdrs and "X-Originating-IP" not in hdrs
    pdf = email_to_pdf(parsed, tmp_path / "print.pdf")
    doc = PdfDocument(pdf)
    doc.extract()
    dets = doc.analyze(RULES)
    out = tmp_path / "o.pdf"
    rep = doc.apply(dets, out)
    t = " ".join(p.get_text() for p in fitz.open(str(out)))
    for v in (SSN, PHONE, EMAIL, "Delgado", "Oakridge", "jbrooks@brookslaw.com", "Venkataraman"):
        assert v not in t, v
    assert "My claim" in t and rep["verification"]["passed"]
    assert "203.0.113.77" not in t
