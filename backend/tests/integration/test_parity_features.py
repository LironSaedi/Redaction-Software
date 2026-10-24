"""CaseGuard-parity features, tested on real output files: picture classification and whole-picture redaction,
reviewer shapes (ellipse / polygon) with per-redaction effects and labels, margins, review copies, the confidence
report, video full-frame effects / enhancement / burned-in captions, and local translation."""
from __future__ import annotations

import io
import shutil
import zipfile
from pathlib import Path

import cv2
import numpy as np
import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw, ImageFont

from obscura.detection.text.engine import AnalyzeOptions
from obscura.redact.model import Box, Detection, ExportOptions
from obscura.redact.pdf import PdfDocument

FIX = Path(__file__).resolve().parent.parent / "fixtures"
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
RULES = AnalyzeOptions(use_spacy=False, use_gliner=False)


def _png(img: Image.Image) -> bytes:
    b = io.BytesIO()
    img.save(b, "PNG")
    return b.getvalue()


def _id_card() -> Image.Image:
    face = Image.open(FIX / "images" / "lena.jpg").resize((220, 260))
    card = Image.new("RGB", (856, 540), (220, 235, 250))
    card.paste(face, (40, 140))
    d = ImageDraw.Draw(card)
    f = ImageFont.truetype(FONT, 34)
    for i, t in enumerate(["OHIO DRIVER LICENSE", "DL TR418822", "DOB 04/17/1986", "DELGADO RUIZ", "MARIA ELENA"]):
        d.text((300, 60 + i * 70), t, fill="black", font=f)
    return card


def _signature() -> Image.Image:
    sig = Image.new("RGB", (900, 260), "white")
    d = ImageDraw.Draw(sig)
    d.line([(40 + i * 8, 150 + int(40 * np.sin(i / 3.0))) for i in range(100)], fill="black", width=4)
    return sig


def _chart() -> Image.Image:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(6, 3), dpi=150)
    ax.bar(["Mar", "Apr", "May"], [2310, 1845, 1960])
    ax.set_title("Medical expenses by month")
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return Image.open(buf)


# ---------------------------------------------------------------------------------------------------
# pictures in PDFs
# ---------------------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def exhibit_pdf(tmp_path_factory):
    d = tmp_path_factory.mktemp("exhibits")
    src = d / "exhibits.pdf"
    doc = fitz.open()
    p = doc.new_page(width=612, height=792)
    p.insert_text((50, 50), "Exhibit B - claimant's photographs and documents", fontsize=11)
    p.insert_image(fitz.Rect(50, 70, 290, 222), stream=_png(_id_card()))
    p.insert_image(fitz.Rect(320, 70, 560, 250), stream=(FIX / "images" / "zidane.jpg").read_bytes())
    p.insert_image(fitz.Rect(50, 300, 330, 380), stream=_png(_signature()))
    p.insert_image(fitz.Rect(50, 420, 450, 620), stream=_png(_chart()))
    doc.save(str(src))
    from obscura.processing import _picture_regions
    pd = PdfDocument(src)
    pd.extract()
    return src, _picture_regions(pd)


def test_every_picture_classified(exhibit_pdf):
    _, regions = exhibit_pdf
    labels = sorted(r.entity for r in regions)
    assert labels == sorted(["IMAGE_ID_CARD", "IMAGE_PHOTO_PERSON", "IMAGE_SIGNATURE", "IMAGE_CHART"]), labels


def test_id_cards_and_signatures_preselected(exhibit_pdf):
    _, regions = exhibit_pdf
    sel = {r.entity: r.selected for r in regions}
    assert sel["IMAGE_ID_CARD"] and sel["IMAGE_SIGNATURE"]
    assert not sel["IMAGE_CHART"] and not sel["IMAGE_PHOTO_PERSON"]


def test_region_detections_are_value_free(exhibit_pdf):
    _, regions = exhibit_pdf
    assert all(r.text == "" and r.kind == "region" and r.meta.get("whole_picture") for r in regions)


def test_whole_picture_redaction_destroys_pixels(exhibit_pdf, tmp_path):
    src, regions = exhibit_pdf
    for r in regions:
        r.selected = True  # redact every picture whole
    out = tmp_path / "out.pdf"
    PdfDocument(src).apply(regions, out, ExportOptions(verify=False))
    page = fitz.open(str(out))[0]
    for r in regions:
        b = r.boxes[0]
        pix = page.get_pixmap(clip=fitz.Rect(b.x0 + 2, b.y0 + 2, b.x1 - 2, b.y1 - 2), dpi=72)
        arr = np.frombuffer(pix.samples, np.uint8)
        assert arr.mean() < 20, (r.entity, arr.mean())  # solid black: nothing of the picture left
    assert "Exhibit B" in page.get_text()


# ---------------------------------------------------------------------------------------------------
# pictures in Word
# ---------------------------------------------------------------------------------------------------

def test_docx_photo_classified_and_redacted_whole(tmp_path):
    import docx
    from docx.shared import Inches

    from obscura.redact.office import OoxmlDocument
    doc = docx.Document()
    doc.add_paragraph("Scene photo attached.")
    doc.add_picture(str(FIX / "images" / "bus.jpg"), width=Inches(4))
    src, dst = tmp_path / "p.docx", tmp_path / "p_red.docx"
    doc.save(str(src))
    od = OoxmlDocument(src)
    dets = od.analyze(RULES)
    pics = [d for d in dets if d.kind == "region"]
    assert pics and pics[0].entity == "IMAGE_PHOTO_PERSON"
    pics[0].selected = True
    od.apply(dets, dst)
    z = zipfile.ZipFile(dst)
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    arr = np.asarray(Image.open(io.BytesIO(z.read(media[0]))).convert("L"))
    assert arr.mean() < 15


# ---------------------------------------------------------------------------------------------------
# shapes, labels, colours, margins, review copies
# ---------------------------------------------------------------------------------------------------

@pytest.fixture()
def words_pdf(tmp_path):
    src = tmp_path / "w.pdf"
    doc = fitz.open()
    p = doc.new_page(width=612, height=792)
    p.insert_text((100, 300), "CENTERWORD", fontsize=12)
    p.insert_text((100, 400), "OUTSIDEWORD", fontsize=12)
    doc.save(str(src))
    return src


def _det(boxes, **meta):
    return Detection(entity="CUSTOM", text="", score=1.0, source="manual:box", kind="manual",
                     boxes=[Box(0, *b) for b in boxes], meta=meta)


def test_polygon_redaction_removes_text_inside(words_pdf, tmp_path):
    out = tmp_path / "o.pdf"
    tri = [[90, 280], [220, 280], [155, 315]]
    PdfDocument(words_pdf).apply([_det([(90, 280, 220, 315)], shape="polygon", points=tri)], out,
                                 ExportOptions(verify=False))
    txt = fitz.open(str(out))[0].get_text()
    assert "CENTERWORD" not in txt and "OUTSIDEWORD" in txt


def test_ellipse_redaction_removes_text_inside(words_pdf, tmp_path):
    out = tmp_path / "o.pdf"
    PdfDocument(words_pdf).apply([_det([(80, 270, 240, 320)], shape="ellipse")], out, ExportOptions(verify=False))
    txt = fitz.open(str(out))[0].get_text()
    assert "CENTERWORD" not in txt and "OUTSIDEWORD" in txt


def test_custom_label_and_colour(words_pdf, tmp_path):
    out = tmp_path / "o.pdf"
    PdfDocument(words_pdf).apply([_det([(95, 285, 190, 305)], label_text="HIPAA", fill=[0.1, 0.3, 0.8])], out,
                                 ExportOptions(verify=False))
    page = fitz.open(str(out))[0]
    assert "HIPAA" in page.get_text() and "CENTERWORD" not in page.get_text()
    pix = page.get_pixmap(clip=fitz.Rect(97, 287, 105, 290), dpi=72)
    r, g, b = np.frombuffer(pix.samples, np.uint8).reshape(-1, 3).mean(axis=0)
    assert b > r + 60  # blue box, not black


def test_padding_widens_the_redaction(words_pdf, tmp_path):
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    box = [(150, 288, 160, 303)]  # a box over the middle of CENTERWORD
    PdfDocument(words_pdf).apply([_det(box)], a, ExportOptions(verify=False))
    PdfDocument(words_pdf).apply([_det(box)], b, ExportOptions(verify=False, padding=40))
    left_a = fitz.open(str(a))[0].get_text().replace("\n", "")
    left_b = fitz.open(str(b))[0].get_text().replace("\n", "")
    assert len(left_b) < len(left_a)


def test_marked_copy_keeps_content_and_is_never_verified(words_pdf, tmp_path):
    out = tmp_path / "review.pdf"
    rep = PdfDocument(words_pdf).apply([_det([(95, 285, 190, 305)])], out, ExportOptions(marked_copy=True))
    txt = fitz.open(str(out))[0].get_text()
    assert "CENTERWORD" in txt and "REVIEW COPY" in txt
    assert rep["marked_copy"] and not rep["verification"]["passed"]


# ---------------------------------------------------------------------------------------------------
# images: polygon effects, enhancement
# ---------------------------------------------------------------------------------------------------

def test_image_polygon_effect_only_inside(tmp_path):
    from obscura.redact.image import ImageDocument
    src = tmp_path / "i.png"
    Image.open(FIX / "images" / "bus.jpg").convert("RGB").save(src)
    doc = ImageDocument(src)
    pts = [[100, 100], [300, 100], [200, 300]]
    det = _det([(100, 100, 300, 300)], shape="polygon", points=pts, effect="blackout")
    out = tmp_path / "o.png"
    doc.apply([det], out, verify=False)
    a = np.asarray(Image.open(out).convert("RGB"))
    assert a[150, 200].sum() == 0          # inside the triangle
    assert a[290, 110].sum() > 0           # inside the box but outside the triangle: untouched


def test_enhancement_brightens(tmp_path):
    from obscura.redact.effects import enhance_frame
    img = (np.random.default_rng(0).random((60, 80, 3)) * 60).astype(np.uint8)
    assert enhance_frame(img, brightness=0.5, contrast=1.2).mean() > img.mean() + 40
    assert enhance_frame(img, sharpen=1.0).shape == img.shape


# ---------------------------------------------------------------------------------------------------
# video: full-frame effect, per-track effect, burned-in captions
# ---------------------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def small_video(tmp_path_factory):
    d = tmp_path_factory.mktemp("vid")
    path = d / "v.mp4"
    shutil.copy(FIX / "media" / "moving_faces.mp4", path)
    return path


def _frame(path, n):
    cap = cv2.VideoCapture(str(path))
    cap.set(cv2.CAP_PROP_POS_FRAMES, n)
    ok, img = cap.read()
    cap.release()
    assert ok
    return img


def test_full_frame_effect_for_a_range(small_video, tmp_path):
    from obscura.redact.video import VideoDocument
    v = VideoDocument(small_video)
    v.tracks = []
    t = v.add_manual_track(5, (0, 0, 0, 0), mode="full_frame", end_frame=15, effect="blackout")
    assert t.label == "full_frame" and len(t.boxes) == 11
    out = tmp_path / "o.mp4"
    v.render(out, verify=False, audio=None)
    assert _frame(out, 10).mean() < 8               # inside the range: black frame
    assert _frame(out, 30).mean() > 20              # outside: picture kept


def test_burned_captions_drawn(small_video, tmp_path):
    from obscura.redact.video import VideoDocument
    v = VideoDocument(small_video)
    v.tracks = []
    out, plain = tmp_path / "c.mp4", tmp_path / "p.mp4"
    v.render(out, verify=False, audio=None, captions=[(0.0, 5.0, "my name is [REDACTED]")])
    v.render(plain, verify=False, audio=None)
    a, b = _frame(out, 3), _frame(plain, 3)
    h = a.shape[0]
    diff = np.abs(a[int(h * 0.75):].astype(int) - b[int(h * 0.75):].astype(int)).mean()
    assert diff > 3  # caption box and text at the bottom


def test_track_effect_survives_round_trip():
    from obscura.detection.vision.tracker import Track
    t = Track(1, "manual", effect="pixelate", shape="ellipse")
    t.boxes = {0: (1, 2, 3, 4)}
    back = Track.from_dict(t.to_dict())
    assert back.effect == "pixelate" and back.shape == "ellipse"


# ---------------------------------------------------------------------------------------------------
# translation (local M2M100)
# ---------------------------------------------------------------------------------------------------

def _translation_available():
    from obscura.detection.audio.translate import available
    return available()


@pytest.mark.skipif(not _translation_available(), reason="translation model not installed")
@pytest.mark.parametrize("target", ["es", "fr", "de", "vi", "ar"])
def test_translation_keeps_redaction_markers(target):
    from obscura.detection.audio.translate import PLACEHOLDER, translate
    out = translate([f"My name is {PLACEHOLDER} and the truck hit my car at the corner."], "en", target)[0]
    assert PLACEHOLDER in out and len(out) > 20


@pytest.mark.skipif(not _translation_available(), reason="translation model not installed")
def test_translation_into_english():
    from obscura.detection.audio.translate import translate
    out = translate(["Me duele mucho la espalda desde el accidente."], "es", "en")[0].lower()
    assert "back" in out and "accident" in out


# ---------------------------------------------------------------------------------------------------
# mailboxes
# ---------------------------------------------------------------------------------------------------

def test_mbox_split_into_messages(tmp_path):
    import mailbox
    from email.message import EmailMessage

    from obscura.processing import kind_for, split_mailbox
    box = mailbox.mbox(str(tmp_path / "client.mbox"))
    for i in range(4):
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = "Maria Delgado <maria.d@example.com>", "atty@example.com", f"Claim {i}"
        m.set_content(f"My SSN is 523-71-992{i}.")
        box.add(m)
    box.flush()
    out = tmp_path / "out"
    out.mkdir()
    msgs = split_mailbox(tmp_path / "client.mbox", "client.mbox", out)
    assert len(msgs) == 4 and all(p.read_bytes().startswith(b"From") or b"Subject" in p.read_bytes() for p in msgs)
    assert kind_for("x.pst")[0] == "mailbox" and kind_for("x.mbox")[0] == "mailbox"


def test_pst_without_readpst_explains_how(tmp_path, monkeypatch):
    import shutil as sh

    from obscura.processing import split_mailbox
    monkeypatch.setattr(sh, "which", lambda name: None)
    (tmp_path / "a.pst").write_bytes(b"!BDN")
    with pytest.raises(ValueError, match="readpst"):
        split_mailbox(tmp_path / "a.pst", "a.pst", tmp_path)


# ---------------------------------------------------------------------------------------------------
# speaker identification
# ---------------------------------------------------------------------------------------------------

def _speaker_model():
    from obscura.detection.audio.speakers import available
    return available()


def test_fbank_shape_and_normalisation():
    from obscura.detection.audio.speakers import fbank
    f = fbank(np.random.default_rng(0).normal(0, 0.1, 16000).astype(np.float32))
    assert f.shape == (98, 80) and abs(float(f.mean())) < 1e-3


@pytest.mark.skipif(not _speaker_model(), reason="speaker model not installed")
def test_two_voices_get_two_speakers():
    from obscura.detection.audio.speakers import diarize
    from obscura.detection.audio.transcribe import decode_16k
    a = decode_16k(FIX / "media" / "speech.aiff")
    rng = np.random.default_rng(1)
    # the same voice twice, then a different "voice" (pitch-shifted by resampling) twice
    other = np.interp(np.arange(0, len(a), 1.35), np.arange(len(a)), a).astype(np.float32)
    seg = 16000 * 2
    audio = np.concatenate([a[:seg], a[seg:2 * seg], other[:seg], other[seg:2 * seg]])
    segs = [{"start": i * 2.0, "end": i * 2.0 + 2.0} for i in range(4)]
    labels = diarize(audio + rng.normal(0, 1e-4, len(audio)).astype(np.float32), segs)
    assert labels[0] == labels[1] and labels[2] == labels[3] and labels[0] != labels[2], labels


# ---------------------------------------------------------------------------------------------------
# concurrency: several background jobs using the models at once must not crash the process
# ---------------------------------------------------------------------------------------------------

def test_models_are_safe_from_parallel_jobs():
    from concurrent.futures import ThreadPoolExecutor

    from obscura.detection.text.engine import AnalyzeOptions, TextAnalyzer
    from obscura.detection.vision.detectors import get_detector
    img = cv2.imread(str(FIX / "images" / "zidane.jpg"))
    an = TextAnalyzer()

    def text_job(i):
        return len(an.analyze(f"Patient Maria Delgado {i} was seen by Dr. Okafor at 4100 Ridgeline Pkwy.", AnalyzeOptions()))

    def vision_job(_):
        return len(get_detector().detect(img, {"face"}))

    with ThreadPoolExecutor(max_workers=6) as ex:
        res = list(ex.map(lambda k: text_job(k) if k % 2 else vision_job(k), range(12)))
    assert all(r > 0 for r in res)
